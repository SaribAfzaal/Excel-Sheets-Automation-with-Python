#!/usr/bin/env python3
"""Merge, clean, format and summarise Excel/CSV files from a folder.

Usage:
    python excel_automation.py --input-folder sample_input --output merged.xlsx \
        --cf-column Sales --threshold 1000
    python excel_automation.py --config config.ini

Command-line arguments override values in the config file.
"""
from __future__ import annotations

import argparse
import configparser
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SUPPORTED = {".xlsx", ".xlsm", ".xls", ".csv"}
SOURCE_COL = "Source File"
FLAG_COL = "Missing Values"

DEFAULTS = {
    "input_folder": "sample_input",
    "output_file": "merged_output.xlsx",
    "missing_strategy": "flag",
    "fill_value": "N/A",
    "numeric_fill_value": "0",
    "cf_column": "",
    "threshold": "",
    "add_source_column": "true",
}


def normalize_columns(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df.columns = [" ".join(str(c).split()).title() for c in df.columns]
    return df


def read_file(path: Path) -> list[pd.DataFrame]:
    if path.suffix.lower() == ".csv":
        frames = {path.name: pd.read_csv(path)}
    else:
        sheets = pd.read_excel(path, sheet_name=None)
        frames = {
            (path.name if len(sheets) == 1 else f"{path.name}:{name}"): df
            for name, df in sheets.items()
        }
    out = []
    for label, df in frames.items():
        df = df.dropna(how="all").dropna(axis=1, how="all")
        if df.empty:
            continue
        df = normalize_columns(df)
        df[SOURCE_COL] = label
        out.append(df)
    return out


def load_and_merge(folder: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Return (merged data, mask of cells that are NaN only because the source
    file never had that column)."""
    files = sorted(p for p in folder.iterdir()
                   if p.suffix.lower() in SUPPORTED and not p.name.startswith(("~$", ".")))
    if not files:
        raise SystemExit(f"No Excel/CSV files found in {folder}")
    frames: list[pd.DataFrame] = []
    for f in files:
        try:
            frames.extend(read_file(f))
            print(f"  loaded {f.name}")
        except Exception as exc:
            print(f"  WARNING: skipped {f.name}: {exc}")
    if not frames:
        raise SystemExit("No readable data found.")
    merged = pd.concat(frames, ignore_index=True, sort=False)
    cols = [c for c in merged.columns if c != SOURCE_COL] + [SOURCE_COL]
    merged = merged[cols]
    absent = pd.concat([pd.DataFrame(False, index=f.index, columns=f.columns) for f in frames],
                       ignore_index=True, sort=False).reindex(columns=cols)
    return merged, absent.fillna(True).astype(bool)


def coerce_numeric(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    for col in df.select_dtypes(include=["object", "string"]).columns:
        if col == SOURCE_COL:
            continue
        cleaned = df[col].astype(str).str.replace(r"[,$\s]", "", regex=True)
        converted = pd.to_numeric(cleaned.where(df[col].notna()), errors="coerce")
        if df[col].notna().sum() and converted.notna().sum() >= 0.8 * df[col].notna().sum():
            df[col] = converted
    return df


def handle_missing(df: pd.DataFrame, absent: pd.DataFrame, strategy: str,
                   fill_value: str, numeric_fill: str) -> tuple[pd.DataFrame, pd.DataFrame]:
    """`absent` marks columns a source file never had; those are not treated as
    genuinely missing for drop/flag, but are still filled by 'fill'."""
    data_cols = [c for c in df.columns if c != SOURCE_COL]
    real_missing = df[data_cols].isna() & ~absent[data_cols]
    if strategy == "drop":
        keep = ~real_missing.any(axis=1)
        return df[keep].reset_index(drop=True), absent[keep].reset_index(drop=True)
    df = df.copy()
    if strategy == "fill":
        for col in data_cols:
            if pd.api.types.is_numeric_dtype(df[col]):
                df[col] = df[col].fillna(float(numeric_fill))
            else:
                df[col] = df[col].fillna(fill_value)
        return df, absent
    if strategy == "flag":
        df[FLAG_COL] = real_missing.apply(lambda r: ", ".join(r.index[r]), axis=1)
        return df, absent
    raise SystemExit(f"Unknown missing strategy: {strategy}")


def clean(df: pd.DataFrame, absent: pd.DataFrame, cfg: dict) -> tuple[pd.DataFrame, dict]:
    stats = {"rows_merged": len(df)}
    df = coerce_numeric(df)
    data_cols = [c for c in df.columns if c != SOURCE_COL]
    before = len(df)
    dup = df.duplicated(subset=data_cols, keep="first")
    df, absent = df[~dup].reset_index(drop=True), absent[~dup].reset_index(drop=True)
    stats["duplicates_removed"] = before - len(df)

    before = len(df)
    df, absent = handle_missing(df, absent, cfg["missing_strategy"], cfg["fill_value"],
                                cfg["numeric_fill_value"])
    stats["rows_dropped_missing"] = before - len(df)
    stats["rows_final"] = len(df)
    if cfg["add_source_column"].lower() not in ("true", "1", "yes"):
        df = df.drop(columns=[SOURCE_COL])
    return df, stats


def build_summary(df: pd.DataFrame) -> pd.DataFrame:
    num = df.select_dtypes(include="number")
    num = num[[c for c in num.columns if not c.lower().endswith("id")]]
    if num.empty:
        return pd.DataFrame({"Info": ["No numeric columns found"]})
    return pd.DataFrame({
        "Column": num.columns,
        "Count": num.count().values,
        "Total": num.sum().values,
        "Average": num.mean().round(2).values,
        "Min": num.min().values,
        "Max": num.max().values,
    })


HEADER_FILL = PatternFill("solid", start_color="1F4E78")
BAND_FILL = PatternFill("solid", start_color="DDEBF7")


def format_sheet(ws, header_color=True) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=2):
        if row[0].row % 2 == 1:
            for cell in row:
                cell.fill = BAND_FILL
    for idx, col in enumerate(ws.columns, start=1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(idx)].width = min(width + 3, 60)
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def add_conditional_formatting(ws, df: pd.DataFrame, column: str, threshold: float) -> None:
    match = next((c for c in df.columns if c.lower() == column.lower()), None)
    if match is None:
        print(f"  WARNING: conditional-format column '{column}' not found. "
              f"Available: {', '.join(df.columns)}")
        return
    if not pd.api.types.is_numeric_dtype(df[match]):
        print(f"  WARNING: column '{match}' is not numeric; skipping conditional formatting")
        return
    letter = get_column_letter(list(df.columns).index(match) + 1)
    rng = f"{letter}2:{letter}{ws.max_row}"
    green = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    ws.conditional_formatting.add(rng, CellIsRule(
        operator="greaterThan", formula=[str(threshold)], fill=green,
        font=Font(color="006100", bold=True)))
    ws.conditional_formatting.add(rng, CellIsRule(
        operator="lessThan", formula=[str(threshold)], fill=red,
        font=Font(color="9C0006", bold=True)))
    print(f"  conditional formatting on '{match}': green > {threshold}, red < {threshold}")


def write_output(df, summary, stats, cfg) -> None:
    out = Path(cfg["output_file"])
    stats_df = pd.DataFrame(
        {"Metric": [k.replace("_", " ").title() for k in stats],
         "Value": list(stats.values())})
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name="Data", index=False)
        summary.to_excel(xw, sheet_name="Summary", index=False)
        stats_df.to_excel(xw, sheet_name="Summary", index=False,
                          startrow=len(summary) + 3)

    wb = load_workbook(out)
    ws = wb["Data"]
    format_sheet(ws)
    if cfg["cf_column"] and cfg["threshold"] != "":
        add_conditional_formatting(ws, df, cfg["cf_column"], float(cfg["threshold"]))
    ss = wb["Summary"]
    format_sheet(ss)
    for cell in ss[len(summary) + 4]:
        if cell.value is not None:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = HEADER_FILL
    for row in ss.iter_rows(min_row=2, max_row=len(summary) + 1):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "#,##0.00"
    wb.save(out)


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--config", help="path to an .ini config file")
    p.add_argument("--input-folder", dest="input_folder")
    p.add_argument("--output", dest="output_file")
    p.add_argument("--missing", dest="missing_strategy", choices=["drop", "fill", "flag"],
                   help="how to handle missing values (default: flag)")
    p.add_argument("--fill-value", dest="fill_value", help="text fill value (default: N/A)")
    p.add_argument("--numeric-fill-value", dest="numeric_fill_value",
                   help="numeric fill value (default: 0)")
    p.add_argument("--cf-column", dest="cf_column",
                   help="column for conditional formatting")
    p.add_argument("--threshold", help="threshold for conditional formatting")
    p.add_argument("--no-source-column", action="store_true",
                   help="omit the 'Source File' column")
    return p.parse_args()


def build_config(args: argparse.Namespace) -> dict:
    cfg = dict(DEFAULTS)
    if args.config:
        cp = configparser.ConfigParser()
        if not cp.read(args.config):
            raise SystemExit(f"Config file not found: {args.config}")
        cfg.update({k: v for k, v in cp["settings"].items()})
    for key in DEFAULTS:
        val = getattr(args, key, None)
        if val is not None:
            cfg[key] = val
    if args.no_source_column:
        cfg["add_source_column"] = "false"
    return cfg


def main() -> None:
    cfg = build_config(parse_args())
    folder = Path(cfg["input_folder"])
    if not folder.is_dir():
        raise SystemExit(f"Input folder not found: {folder}")
    print(f"Reading files from {folder}")
    merged, absent = load_and_merge(folder)
    df, stats = clean(merged, absent, cfg)
    summary = build_summary(df)
    write_output(df, summary, stats, cfg)
    print(f"\nDone -> {cfg['output_file']}")
    for k, v in stats.items():
        print(f"  {k.replace('_', ' ')}: {v}")


if __name__ == "__main__":
    sys.exit(main())
